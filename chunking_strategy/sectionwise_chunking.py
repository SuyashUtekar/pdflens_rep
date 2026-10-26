"""Create section-aware text chunks from a Word document.

Run from the project root with:
    uv run python chunking_strategy/sectionwise_chunking.py path/to/file.docx

The output is JSON beside the source DOCX unless ``--output`` is supplied.
"""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

try:
    from docx import Document
    from docx.document import Document as DocumentType
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError as exc:  # pragma: no cover - exercised by the CLI error path
    raise SystemExit(
        "python-docx is required. Install project dependencies with `uv sync`, "
        "then run this script with `uv run python`."
    ) from exc


@dataclass
class Section:
    """A heading and the content that belongs directly to it."""

    title: str
    level: int
    path: list[str]
    blocks: list[str] = field(default_factory=list)


def _paragraph_has_image(paragraph: Paragraph) -> bool:
    # python-docx's XPath helper does not support namespace prefixes on every
    # supported release, so inspect the already parsed paragraph XML directly.
    xml = paragraph._p.xml
    return "<w:drawing" in xml or "<w:pict" in xml


def _table_to_text(table: Table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
        if any(cells):
            rows.append(" ; ".join(cells))
    return "\n".join(rows)


def _iter_blocks(document: DocumentType) -> Iterable[tuple[str, str, int | None]]:
    """Yield document-order blocks as (kind, text, heading_level)."""
    image_number = 0
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            style_name = block.style.name if block.style is not None else ""
            if style_name == "Title":
                if text:
                    yield "title", text, None
                continue

            match = re.fullmatch(r"Heading\s+(\d+)", style_name)
            if match:
                if text:
                    yield "heading", text, int(match.group(1))
                continue

            if text:
                yield "text", text, None
            if _paragraph_has_image(block):
                image_number += 1
                yield "image", f"[Embedded image {image_number}]", None
        elif isinstance(block, Table):
            text = _table_to_text(block)
            if text:
                yield "table", text, None


def _split_words(text: str, max_words: int, overlap_words: int) -> list[str]:
    words = text.split()
    if len(words) <= max_words:
        return [text] if text.strip() else []

    chunks: list[str] = []
    start = 0
    while start < len(words):
        end = min(start + max_words, len(words))
        chunks.append(" ".join(words[start:end]))
        if end == len(words):
            break
        start = max(end - overlap_words, start + 1)
    return chunks


def chunk_docx(
    source: Path,
    *,
    max_words: int = 800,
    overlap_words: int = 80,
) -> dict[str, object]:
    """Read a DOCX and return chunks grouped by its Word heading hierarchy."""
    if max_words < 1:
        raise ValueError("max_words must be at least 1")
    if overlap_words < 0 or overlap_words >= max_words:
        raise ValueError("overlap_words must be between 0 and max_words - 1")

    document = Document(str(source))
    title = ""
    sections: list[Section] = []
    active: Section | None = None
    heading_stack: list[tuple[int, str]] = []
    preamble: list[str] = []

    def flush() -> None:
        nonlocal active
        if active is not None:
            sections.append(active)
            active = None

    for kind, text, level in _iter_blocks(document):
        if kind == "title":
            title = text
            continue
        if kind == "heading":
            flush()
            assert level is not None
            while heading_stack and heading_stack[-1][0] >= level:
                heading_stack.pop()
            heading_stack.append((level, text))
            active = Section(
                title=text,
                level=level,
                path=[heading for _, heading in heading_stack],
            )
            continue
        if active is None:
            if heading_stack:
                active = Section(
                    title=heading_stack[-1][1],
                    level=heading_stack[-1][0],
                    path=[heading for _, heading in heading_stack],
                )
            else:
                preamble.append(text)
                continue
        active.blocks.append(text)
    flush()

    if preamble:
        sections.insert(0, Section(title="Introduction", level=0, path=[], blocks=preamble))

    chunks: list[dict[str, object]] = []
    for section in sections:
        section_text = "\n\n".join(section.blocks).strip()
        if not section_text:
            continue
        for part_number, part in enumerate(
            _split_words(section_text, max_words=max_words, overlap_words=overlap_words),
            start=1,
        ):
            chunk_parts = [*section.path, part] if section.path else [title, part]
            chunk_text = " | ".join(value for value in chunk_parts if value)
            chunks.append(
                {
                    "chunk_id": f"chunk-{len(chunks) + 1:06d}",
                    "document": source.name,
                    "document_title": title or source.stem,
                    "section_title": section.title,
                    "section_level": section.level,
                    "section_path": section.path,
                    "part": part_number,
                    "text": chunk_text,
                    "word_count": len(chunk_text.split()),
                }
            )

    return {
        "source_docx": str(source.resolve()),
        "document_title": title or source.stem,
        "chunk_count": len(chunks),
        "max_words": max_words,
        "overlap_words": overlap_words,
        "chunks": chunks,
    }


def save_chunks_xlsx(chunks: list[dict[str, object]], destination: Path) -> None:
    """Write the requested two-column Excel view of the generated chunks."""
    try:
        from openpyxl import Workbook as ExcelWorkbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:  # pragma: no cover - exercised by the CLI error path
        raise SystemExit(
            "openpyxl is required for Excel output. Install project dependencies with `uv sync`."
        ) from exc

    workbook = ExcelWorkbook()
    sheet = workbook.active
    sheet.title = "Chunks"
    sheet.append(["chunk_id", "chunk_text"])
    for chunk in chunks:
        sheet.append([chunk["chunk_id"], chunk["text"]])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 100
    sheet.row_dimensions[1].height = 24
    for row_number in range(2, sheet.max_row + 1):
        text = str(sheet.cell(row=row_number, column=2).value or "")
        wrapped_lines = sum(max(1, (len(line) + 98) // 99) for line in text.splitlines())
        sheet.row_dimensions[row_number].height = max(30, min(wrapped_lines * 15 + 8, 600))
        for cell in sheet[row_number]:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Split a DOCX into chunks grouped by Word heading sections."
    )
    parser.add_argument("docx", type=Path, help="Path to the input DOCX file")
    parser.add_argument(
        "--output", "-o", type=Path, help="Output JSON path (defaults beside the DOCX)"
    )
    parser.add_argument("--max-words", type=int, default=800, help="Maximum words per chunk")
    parser.add_argument(
        "--overlap-words", type=int, default=80, help="Words repeated between split chunks"
    )
    parser.add_argument(
        "--excel-output",
        type=Path,
        help="Excel output path (defaults beside the JSON output)",
    )
    args = parser.parse_args()

    source = args.docx.expanduser().resolve()
    if not source.is_file():
        parser.error(f"DOCX file does not exist: {source}")
    if source.suffix.lower() != ".docx":
        parser.error(f"Expected a .docx file: {source}")

    output = args.output.expanduser().resolve() if args.output else source.with_name(
        f"{source.stem}_chunks.json"
    )
    result = chunk_docx(source, max_words=args.max_words, overlap_words=args.overlap_words)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    excel_output = args.excel_output.expanduser().resolve() if args.excel_output else output.with_suffix(
        ".xlsx"
    )
    save_chunks_xlsx(result["chunks"], excel_output)
    print(f"Source: {source}")
    print(f"Chunks: {result['chunk_count']}")
    print(f"Saved: {output}")
    print(f"Saved: {excel_output}")


if __name__ == "__main__":
    main()
