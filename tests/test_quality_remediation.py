from __future__ import annotations

import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from app.artifacts.package import build_package
from app.parsers.base import ParsedAsset
from app.parsers.docling_adapter import DoclingAdapter
from app.schemas.models import (
    BoundingBox,
    DocumentInfo,
    Element,
    ExtractionField,
    Metadata,
    Page,
    ProcessingInfo,
)


def make_page(page_number: int = 1) -> Page:
    return Page(page_number=page_number, physical_page_index=page_number,
                width_pt=612, height_pt=792)


class QualityRemediationTests(unittest.TestCase):
    def test_pypdf_fallback_fills_empty_and_partially_extracted_pages(self) -> None:
        empty_page = make_page(1)
        populated_page = make_page(2)
        populated_page.elements.append(Element(
            element_id="p0002-e0001", type=ExtractionField.BODY_TEXT,
            text="Docling text", reading_order=0,
        ))

        DoclingAdapter._apply_empty_page_fallback(
            [empty_page, populated_page], ["Recovered source text", "Source baseline"]
        )

        self.assertEqual(len(empty_page.elements), 1)
        self.assertEqual(empty_page.elements[0].text, "Recovered source text")
        self.assertEqual(empty_page.elements[0].attributes["source"], "pypdf_fallback")
        self.assertIn("BODY_TEXT_FALLBACK_USED", empty_page.warnings)
        self.assertEqual(len(populated_page.elements), 2)
        self.assertEqual(populated_page.elements[0].text, "Docling text")
        recovered = populated_page.elements[1]
        self.assertEqual(recovered.text, "Source baseline")
        self.assertEqual(recovered.attributes["fallback_reason"], "unmatched_source_text_span")
        self.assertIn("BODY_TEXT_PARTIAL_FALLBACK_USED", populated_page.warnings)
        self.assertLess(populated_page.source_text_coverage_before_recovery, 0.9)
        self.assertEqual(populated_page.source_text_coverage_ratio, 1.0)

    def test_empty_source_page_is_flagged_for_review(self) -> None:
        page = make_page()
        DoclingAdapter._apply_empty_page_fallback([page], [""])
        self.assertIn("BODY_TEXT_SOURCE_LAYER_EMPTY", page.warnings)

    def test_fallback_removes_text_already_emitted_as_a_typed_element(self) -> None:
        page = make_page()
        page.elements.append(Element(
            element_id="p0001-heading", type=ExtractionField.HEADING,
            text="Important Heading", reading_order=0,
        ))

        DoclingAdapter._apply_empty_page_fallback(
            [page], ["Important\nHeading\nRemaining body text"]
        )

        fallback = next(element for element in page.elements
                        if element.attributes.get("source") == "pypdf_fallback")
        self.assertEqual(fallback.text, "Remaining body text")
        self.assertEqual(fallback.attributes["overlap_removed_char_count"], len("Important Heading"))
        self.assertIn("BODY_TEXT_FALLBACK_OVERLAP_REMOVED", page.warnings)

    def test_doc11_reported_empty_body_pages_get_source_text_fallback(self) -> None:
        from pypdf import PdfReader

        pdf_path = Path(__file__).parents[1] / "rbi_docs" / "doc11.pdf"
        source_pages = PdfReader(pdf_path, strict=False).pages
        source_texts = [source_page.extract_text() or "" for source_page in source_pages]
        reported_empty_pages = {5, 8, 9, 10, 11, 12, 14, 17, 19}
        pages = [make_page(index) for index in range(1, len(source_pages) + 1)]
        for page in pages:
            if page.page_number not in reported_empty_pages:
                page.elements.append(Element(
                    element_id=f"p{page.page_number:04d}-docling",
                    type=ExtractionField.BODY_TEXT, text="Existing Docling text",
                    reading_order=0,
                ))

        DoclingAdapter._apply_empty_page_fallback(pages, source_texts)

        for page_number in reported_empty_pages:
            page = pages[page_number - 1]
            self.assertTrue(any(element.attributes.get("source") == "pypdf_fallback"
                                for element in page.elements))
        for page in pages:
            if page.page_number not in reported_empty_pages:
                self.assertGreaterEqual(len(page.elements), 1)
                if source_texts[page.page_number - 1].strip():
                    self.assertIn("BODY_TEXT_PARTIAL_FALLBACK_USED", page.warnings)

    def test_reported_critical_pages_are_recovered_from_independent_source_text(self) -> None:
        from pypdf import PdfReader

        targets = {
            "doc30.pdf": {1, 4, 5, 11, 12, 18, 23, 30, 32, 36,
                          43, 44, 45, 46, 47, 48, 49, 50, 51, 52},
            "doc11.pdf": {3, 4, 6, 7, 13, 15, 16, 18},
        }
        root = Path(__file__).parents[1] / "rbi_docs"
        for filename, target_pages in targets.items():
            source_texts = [page.extract_text() or "" for page in
                            PdfReader(root / filename, strict=False).pages]
            pages = [make_page(index) for index in range(1, len(source_texts) + 1)]
            for page, source_text in zip(pages, source_texts):
                words = source_text.split()
                simulated_output = " ".join(words[:4]) if page.page_number in target_pages else source_text
                if simulated_output:
                    page.elements.append(Element(
                        element_id=f"p{page.page_number:04d}-simulated-primary",
                        type=ExtractionField.BODY_TEXT,
                        text=simulated_output,
                        reading_order=0,
                    ))

            DoclingAdapter._apply_empty_page_fallback(pages, source_texts)

            for page_number in target_pages:
                page = pages[page_number - 1]
                self.assertLess(page.source_text_coverage_before_recovery, 0.90,
                                f"{filename} page {page_number} was not flagged")
                self.assertGreaterEqual(page.source_text_coverage_ratio, 0.90,
                                        f"{filename} page {page_number} remained incomplete")
                self.assertTrue(page.fallback_used)
                self.assertGreater(page.fallback_chars_added, 0)

    def test_numbered_clause_is_not_footnote_without_lower_page_position(self) -> None:
        top_box = BoundingBox(x0=0.1, y0=0.2, x1=0.8, y1=0.23)
        bottom_box = BoundingBox(x0=0.1, y0=0.9, x1=0.8, y1=0.93)
        self.assertEqual(
            DoclingAdapter._map_label("text", "1. A numbered clause", top_box),
            ExtractionField.BODY_TEXT,
        )
        self.assertEqual(
            DoclingAdapter._map_label("text", "1. A note", bottom_box),
            ExtractionField.FOOTNOTE,
        )

    def test_repeated_page_edge_text_is_classified_without_removal(self) -> None:
        pages = [make_page(1), make_page(2)]
        for index, page in enumerate(pages, start=1):
            page.elements.append(Element(
                element_id=f"p{index:04d}-e0001", type=ExtractionField.BODY_TEXT,
                text="Regulatory Directions", reading_order=0,
                bbox=BoundingBox(x0=0.1, y0=0.03, x1=0.9, y1=0.08),
            ))
        DoclingAdapter._classify_repeated_page_edges(pages)
        self.assertTrue(all(page.elements[0].type == ExtractionField.HEADER for page in pages))
        self.assertTrue(all(page.elements[0].text == "Regulatory Directions" for page in pages))
        self.assertTrue(all(page.elements[0].confidence == 0.8 for page in pages))

    def test_repeating_pdf_edge_text_is_added_but_one_off_title_is_not(self) -> None:
        class FakePage:
            mediabox = SimpleNamespace(height=792, width=612)

            def __init__(self, label: str):
                self.label = label

            def extract_text(self, visitor_text):
                visitor_text(self.label, [1, 0, 0, 1, 72, 740],
                             [1, 0, 0, 1, 0, 0], None, 12)

        class FakeReader:
            def __init__(self, *_args, **_kwargs):
                self.pages = [FakePage("Running Header"), FakePage("Running Header"),
                              FakePage("Unique section title")]

        pages = [make_page(index) for index in range(1, 4)]
        with patch("pypdf.PdfReader", FakeReader):
            DoclingAdapter._append_repeated_pdf_edge_furniture("unused.pdf", pages)

        self.assertTrue(all(any(element.type == ExtractionField.HEADER
                                 and element.text == "Running Header"
                                 for element in page.elements) for page in pages[:2]))
        self.assertFalse(any(element.type == ExtractionField.HEADER
                             and element.text == "Unique section title"
                             for element in pages[2].elements))

    def test_printed_page_labels_are_separate_from_physical_indices(self) -> None:
        fixture = Path(__file__).parents[1] / "docs" / "PDF_Layout_Metadata_Test_Fixture.pdf"
        from pypdf import PdfReader

        source_pages = PdfReader(fixture).pages
        pages = [
            Page(page_number=index, physical_page_index=index,
                 width_pt=float(source_page.mediabox.width),
                 height_pt=float(source_page.mediabox.height))
            for index, source_page in enumerate(source_pages, start=1)
        ]
        DoclingAdapter._append_pdf_page_numbers(
            str(fixture), pages, {ExtractionField.PAGE_NUMBER}
        )
        self.assertEqual([page.printed_page_label for page in pages],
                         [f"Page {index}" for index in range(1, len(pages) + 1)])
        self.assertEqual([page.physical_page_index for page in pages],
                         list(range(1, len(pages) + 1)))

    def test_page_number_annotation_contract_emits_at_most_one_element_per_page(self) -> None:
        class FakePage:
            mediabox = SimpleNamespace(height=792)

            def extract_text(self, visitor_text):
                identity = [1, 0, 0, 1, 0, 0]
                visitor_text("Page 1", identity, [1, 0, 0, 1, 0, 30], None, 10)
                visitor_text("1", identity, [1, 0, 0, 1, 0, 30], None, 10)

        class FakeReader:
            def __init__(self, *_args, **_kwargs):
                self.pages = [FakePage()]

        page = make_page()
        with patch("pypdf.PdfReader", FakeReader):
            DoclingAdapter._append_pdf_page_numbers("unused.pdf", [page],
                                                    {ExtractionField.PAGE_NUMBER})

        number_elements = [element for element in page.elements
                           if element.type == ExtractionField.PAGE_NUMBER]
        self.assertEqual(len(number_elements), 1)
        self.assertEqual(page.printed_page_label, "Page 1")

    def test_internal_pdf_links_keep_page_targets_and_visible_text(self) -> None:
        destination_page = object()

        class FakeAnnotation:
            def get_object(self):
                return {"/A": {"/S": "/GoTo", "/D": [destination_page]},
                        "/Rect": [60, 700, 180, 730]}

        class FakePdfPage:
            mediabox = SimpleNamespace(width=612, height=792)

            def get(self, key, default=None):
                return {"/Annots": [FakeAnnotation()]}.get(key, default)

        class FakeReader:
            pages = [FakePdfPage()]

            def __init__(self, *_args, **_kwargs):
                pass

            def get_page_number(self, page_ref):
                self.asserted_ref = page_ref
                return 1

        page = make_page()
        page.elements.append(Element(
            element_id="index-entry", type=ExtractionField.BODY_TEXT,
            text="Security controls", reading_order=0,
            bbox=BoundingBox(x0=0.09, y0=0.07, x1=0.30, y1=0.12),
        ))
        warnings: list[str] = []
        with patch("pypdf.PdfReader", FakeReader):
            DoclingAdapter._append_internal_pdf_links("unused.pdf", [page], warnings)

        link = next(element for element in page.elements
                    if element.type == ExtractionField.HYPERLINK)
        self.assertEqual(link.target_uri, "#page=2")
        self.assertEqual(link.text, "Security controls")
        self.assertEqual(link.attributes["target_type"], "internal_page")
        self.assertEqual(page.hyperlink_count, 1)

    def test_docling_table_cells_are_preserved_as_structured_data(self) -> None:
        from docling_core.types.doc import TableCell

        cell = TableCell(
            text="Threshold", start_row_offset_idx=0, end_row_offset_idx=1,
            start_col_offset_idx=0, end_col_offset_idx=1, column_header=True,
        )
        table = SimpleNamespace(data=SimpleNamespace(table_cells=[cell]))
        element = Element(element_id="p0001-e0001", type=ExtractionField.TABLE,
                          reading_order=0)
        DoclingAdapter._attach_table_structure(table, element, make_page())
        self.assertEqual(element.attributes["table_status"], "structured")
        self.assertEqual(element.attributes["table_cells"][0]["text"], "Threshold")
        self.assertEqual(element.attributes["table_validation_status"], "needs_review")
        self.assertEqual(element.attributes["table_metrics"]["row_count"], 1)
        self.assertEqual(element.attributes["table_metrics"]["column_count"], 1)

    def test_table_without_cells_is_marked_visual_only(self) -> None:
        table = SimpleNamespace(data=SimpleNamespace(table_cells=[]))
        page = make_page()
        element = Element(element_id="p0001-e0001", type=ExtractionField.TABLE,
                          reading_order=0)
        DoclingAdapter._attach_table_structure(table, element, page)
        self.assertEqual(element.attributes["table_status"], "visual_only")
        self.assertIn("TABLE_VISUAL_ONLY", page.warnings)

    def test_package_contains_source_and_manifest_with_no_html_file(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source.pdf"
            source.write_bytes(b"%PDF-test source bytes")
            output = root / "sample.zip"
            metadata = Metadata(
                document=DocumentInfo(source_filename="sample.pdf", source_sha256="a" * 64,
                                      page_count=0),
                processing=ProcessingInfo(
                    parser_name="docling", parser_version="test", requested_fields=[],
                    extracted_fields=[],
                ),
                pages=[],
            )
            build_package(metadata, output, source)
            with zipfile.ZipFile(output) as archive:
                names = set(archive.namelist())
                self.assertIn("sample.pdf", names)
                self.assertIn("document.docx", names)
                self.assertIn("metadata.json", names)
                self.assertIn("manifest.json", names)
                self.assertFalse(any(name.endswith(".html") for name in names))
                manifest = json.loads(archive.read("manifest.json"))
                self.assertEqual(manifest["metadata_schema_version"], "1.4")
                self.assertEqual({item["path"] for item in manifest["members"]},
                                 names - {"manifest.json"})
                self.assertIsNone(json.loads(archive.read("metadata.json"))
                                  ["artifacts"]["html_path"])

    def test_package_rejects_metadata_referencing_missing_asset(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source.pdf"
            source.write_bytes(b"%PDF-test source bytes")
            output = root / "sample.zip"
            page = make_page()
            page.elements.append(Element(
                element_id="p0001-e0001", type=ExtractionField.IMAGE,
                reading_order=0, artifact_path="images/missing.png",
            ))
            metadata = Metadata(
                document=DocumentInfo(source_filename="sample.pdf", source_sha256="a" * 64,
                                      page_count=1),
                processing=ProcessingInfo(
                    parser_name="docling", parser_version="test", requested_fields=[],
                    extracted_fields=[],
                ),
                pages=[page],
            )
            with self.assertRaisesRegex(ValueError, "missing package artifacts"):
                build_package(metadata, output, source)

    def test_package_deduplicates_identical_image_bytes_but_keeps_occurrences(self) -> None:
        from io import BytesIO
        from PIL import Image

        stream = BytesIO()
        Image.new("RGB", (2, 2), color="white").save(stream, format="PNG")
        png = stream.getvalue()
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source.pdf"
            source.write_bytes(b"%PDF-test source bytes")
            output = root / "sample.zip"
            first, second = make_page(1), make_page(2)
            first.elements.append(Element(element_id="image-one", type=ExtractionField.IMAGE,
                                          reading_order=0))
            second.elements.append(Element(element_id="image-two", type=ExtractionField.IMAGE,
                                           reading_order=0))
            metadata = Metadata(
                document=DocumentInfo(source_filename="sample.pdf", source_sha256="a" * 64,
                                      page_count=2),
                processing=ProcessingInfo(parser_name="docling", parser_version="test",
                                          requested_fields=[ExtractionField.IMAGE],
                                          extracted_fields=[ExtractionField.IMAGE]),
                pages=[first, second],
            )
            assets = [
                ParsedAsset("image-one", 1, ExtractionField.IMAGE, png),
                ParsedAsset("image-two", 2, ExtractionField.IMAGE, png),
            ]

            build_package(metadata, output, source, assets)

            self.assertEqual(first.elements[0].artifact_path, second.elements[0].artifact_path)
            with zipfile.ZipFile(output) as archive:
                image_paths = [name for name in archive.namelist()
                               if name.startswith("images/") and name.endswith(".png")]
                self.assertEqual(len(image_paths), 1)

    def test_body_text_assembly_removes_exact_duplicate_elements(self) -> None:
        from app.main import _assemble_body_text

        page = make_page()
        page.elements.extend([
            Element(element_id="first", type=ExtractionField.BODY_TEXT,
                    text="Repeated paragraph", reading_order=0),
            Element(element_id="duplicate", type=ExtractionField.BODY_TEXT,
                    text="  repeated   paragraph ", reading_order=1),
            Element(element_id="next", type=ExtractionField.BODY_TEXT,
                    text="Next paragraph", reading_order=2),
        ])

        self.assertEqual(_assemble_body_text(page), "Repeated paragraph\nNext paragraph")
        self.assertEqual(page.duplicate_text_char_count, len("  repeated   paragraph "))
        self.assertIn("BODY_TEXT_EXACT_DUPLICATE_REMOVED", page.warnings)

    def test_body_text_assembly_includes_structural_body_text_not_furniture(self) -> None:
        from app.main import _assemble_body_text

        page = make_page()
        page.elements.extend([
            Element(element_id="heading", type=ExtractionField.HEADING,
                    text="Chapter heading", reading_order=0),
            Element(element_id="list", type=ExtractionField.LIST,
                    text="Requirement one", reading_order=1),
            Element(element_id="paragraph", type=ExtractionField.BODY_TEXT,
                    text="Body paragraph", reading_order=2),
            Element(element_id="header", type=ExtractionField.HEADER,
                    text="Running header", reading_order=3),
            Element(element_id="footnote", type=ExtractionField.FOOTNOTE,
                    text="Separate footnote", reading_order=4),
        ])

        self.assertEqual(_assemble_body_text(page),
                         "Chapter heading\nRequirement one\nBody paragraph")

    def test_low_body_coverage_cannot_pass_page_quality_gate(self) -> None:
        from app.main import _assess_page_quality

        page = make_page()
        page.source_text_char_count = 1289
        page.body_text_char_count = 226
        page.body_text_to_source_text_ratio = 226 / 1289
        self.assertEqual(_assess_page_quality(page), "failed")
        self.assertIn("BODY_TEXT_QUALITY_GATE_FAILED", page.warnings)

        table_page = make_page()
        table_page.source_text_char_count = 1289
        table_page.body_text_char_count = 226
        table_page.body_text_to_source_text_ratio = 226 / 1289
        table_page.elements.append(Element(
            element_id="table", type=ExtractionField.TABLE, reading_order=0,
            attributes={"table_validation_status": "needs_review"},
        ))
        self.assertEqual(_assess_page_quality(table_page), "review")

    def test_unicode_replacement_warning_requires_review_even_with_high_coverage(self) -> None:
        from app.main import _assess_page_quality

        page = make_page()
        page.source_text_char_count = 100
        page.body_text_char_count = 98
        page.body_text_to_source_text_ratio = 0.98
        page.warnings.append("UNICODE_SUSPECT")

        self.assertEqual(_assess_page_quality(page), "review")

    def test_docx_uses_source_geometry_native_tables_and_internal_bookmarks(self) -> None:
        from docx import Document as WordDocument

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source.pdf"
            source.write_bytes(b"%PDF-test source bytes")
            output = root / "sample.zip"
            first = Page(page_number=1, width_pt=595.98, height_pt=841.98)
            first.elements.extend([
                Element(element_id="toc-link", type=ExtractionField.HYPERLINK,
                        text="Annex", target_uri="#page=2", reading_order=0),
                Element(element_id="table", type=ExtractionField.TABLE,
                        reading_order=1, attributes={"table_cells": [
                            {"text": "Name", "row_index": 0, "column_index": 0,
                             "is_column_header": True},
                            {"text": "Value", "row_index": 0, "column_index": 1,
                             "is_column_header": True},
                            {"text": "Rate", "row_index": 1, "column_index": 0},
                            {"text": "5%", "row_index": 1, "column_index": 1},
                        ]}),
            ])
            second = Page(page_number=2, width_pt=595.98, height_pt=841.98)
            second.elements.append(Element(
                element_id="target-heading", type=ExtractionField.HEADING,
                text="Annex", reading_order=0,
            ))
            metadata = Metadata(
                document=DocumentInfo(source_filename="sample.pdf", source_sha256="a" * 64,
                                      page_count=2),
                processing=ProcessingInfo(parser_name="docling", parser_version="test",
                                          requested_fields=[], extracted_fields=[]),
                pages=[first, second],
            )
            build_package(metadata, output, source)

            with zipfile.ZipFile(output) as archive:
                document = WordDocument(io.BytesIO(archive.read("document.docx")))
                xml = document.part.element.xml
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(round(document.sections[0].page_width.pt), 596)
            self.assertIn('w:anchor="pdflens_page_2"', xml)
            self.assertIn('w:name="pdflens_page_2"', xml)
            self.assertNotIn("#page=2", "\n".join(p.text for p in document.paragraphs))


if __name__ == "__main__":
    unittest.main()
